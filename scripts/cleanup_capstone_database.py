from __future__ import annotations

import json
import os
import sqlite3
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "attendance.db"
WORKBOOK_PATH = ROOT.parents[1] / "Database for FRAS (1).xlsx"


DEPT_IDS = {
    "DLA": 1,
    "MATH": 2,
    "PHYS": 3,
    "ETYBSM": 4,
    "SOIT": 5,
    "SOMDA": 6,
}


def dept_for_course(code: str) -> int:
    prefix = "".join(ch for ch in code.upper() if ch.isalpha())
    if prefix == "MATH":
        return DEPT_IDS["MATH"]
    if prefix == "PHYS":
        return DEPT_IDS["PHYS"]
    if prefix in {"GED", "FW", "NSTP", "RZL", "SGE"}:
        return DEPT_IDS["DLA"]
    if prefix in {"CS", "IT", "IS", "CIS", "EMC"}:
        return DEPT_IDS["SOIT"]
    return DEPT_IDS["SOIT"]


def units_for_course(code: str, title: str) -> int:
    upper_code = code.upper()
    upper_title = title.upper()
    if upper_code.endswith("L") or upper_code.endswith("-8L") or "LABORATORY" in upper_title:
        return 1
    if upper_code.startswith("FW"):
        return 2
    if "PRACTICUM" in upper_title or "THESIS" in upper_title:
        return 3
    return 3


def load_courses() -> list[tuple[int, str, str, int, int]]:
    workbook = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
    sheet = workbook["Courses"]
    seen: set[str] = set()
    courses: list[tuple[str, str, int, int]] = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        raw_code, raw_title = row[:2]
        if not raw_code or not raw_title:
            continue
        code = str(raw_code).strip()
        title = str(raw_title).strip().title()
        if code.upper() in seen:
            continue
        seen.add(code.upper())
        courses.append((code, title, units_for_course(code, title), dept_for_course(code)))

    extra_courses = [
        ("IT119L", "Research Methods In Information Technology Laboratory", 1, DEPT_IDS["SOIT"]),
        ("CIS101", "Cisco Networking Fundamentals", 3, DEPT_IDS["SOIT"]),
    ]
    for code, title, units, dept_id in extra_courses:
        if code.upper() not in seen:
            seen.add(code.upper())
            courses.append((code, title, units, dept_id))

    return [
        (index, code, title, units, dept_id)
        for index, (code, title, units, dept_id) in enumerate(courses, start=1)
    ]


def set_sequence(cursor: sqlite3.Cursor, table_name: str, value: int) -> None:
    cursor.execute("DELETE FROM sqlite_sequence WHERE name = ?", (table_name,))
    cursor.execute(
        "INSERT INTO sqlite_sequence (name, seq) VALUES (?, ?)",
        (table_name, value),
    )


def cleanup_database() -> dict[str, int]:
    keep_user_emails = {
        "david.miller@mapua.edu.ph",
        "emily.williams@example.com",
        "eric.deveza@mapua.edu.ph",
        "isaac.romance@mapua.edu.ph",
        "jane.smith@mapua.edu.ph",
        "john.johnson@mapua.edu.ph",
        "lisa.wilson@mapua.edu.ph",
        "maria.santos@mapua.edu.ph",
        "michael.brown@mapua.edu.ph",
        "sarah.davis@mapua.edu.ph",
        "admin@mapua.edu.ph",
        "superadmin@fras.com",
    }

    courses = load_courses()

    with sqlite3.connect(DB_PATH) as conn:
        cursor = conn.cursor()
        cursor.execute("PRAGMA foreign_keys = OFF")
        cursor.execute("BEGIN")

        try:
            cursor.execute("DELETE FROM ticket_replies")
            cursor.execute("DELETE FROM support_tickets")
            cursor.execute("DELETE FROM audit_logs")
            cursor.execute("DELETE FROM student_face_embeddings")

            cursor.execute(
                """
                DELETE FROM attendance_logs
                WHERE class_id NOT IN (8, 9)
                   OR date(timestamp) NOT IN ('2026-01-15', '2026-02-05')
                   OR (class_id = 8 AND date(timestamp) != '2026-02-05')
                   OR (class_id = 9 AND date(timestamp) != '2026-01-15')
                """
            )
            cursor.execute(
                """
                UPDATE attendance_logs
                SET notes = 'Marked absent after attendance window closed.'
                WHERE status_id = 2
                """
            )

            cursor.execute("DELETE FROM enrollments WHERE class_id NOT IN (8, 9)")
            cursor.execute("DELETE FROM classes WHERE class_id NOT IN (8, 9)")

            placeholders = ",".join("?" for _ in keep_user_emails)
            cursor.execute(
                f"DELETE FROM users WHERE email NOT IN ({placeholders})",
                tuple(sorted(keep_user_emails)),
            )
            cursor.execute("DELETE FROM instructors WHERE instructor_id NOT BETWEEN 1 AND 10")
            cursor.execute("DELETE FROM admins WHERE admin_id != 1")
            cursor.execute("DELETE FROM it_admins WHERE it_admin_id != 1")
            cursor.execute("DELETE FROM super_admins WHERE super_admin_id != 1")

            cursor.execute("DELETE FROM students WHERE student_id = 151")

            cursor.execute("DELETE FROM rooms WHERE room_id NOT IN (10, 11)")
            cursor.execute(
                """
                UPDATE rooms
                SET room_number = '604',
                    floor_level = 6,
                    campus_id = 1,
                    building_id = 1,
                    room_type_id = 2
                WHERE room_id = 10
                """
            )
            cursor.execute(
                """
                UPDATE rooms
                SET room_number = '318',
                    floor_level = 3,
                    campus_id = 1,
                    building_id = 1,
                    room_type_id = 2
                WHERE room_id = 11
                """
            )
            cursor.executemany(
                """
                INSERT OR REPLACE INTO rooms (
                    room_id, room_number, floor_level, campus_id, building_id, room_type_id
                )
                VALUES (?, ?, ?, 1, 1, 2)
                """,
                [(13, "606", 6), (14, "506", 5)],
            )

            cursor.execute("DELETE FROM courses")
            cursor.executemany(
                """
                INSERT INTO courses (
                    course_id, course_code, course_name, units, dept_id, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
                """,
                courses,
            )

            cursor.execute("SELECT course_id FROM courses WHERE course_code = 'IT119'")
            it119_id = cursor.fetchone()[0]
            cursor.execute("SELECT course_id FROM courses WHERE course_code = 'IT119L'")
            it119l_id = cursor.fetchone()[0]

            cursor.execute(
                """
                UPDATE classes
                SET course_id = ?,
                    room_id = 10,
                    instructor_id = 10,
                    section = 'AM1',
                    day_of_week = 'Thursday',
                    start_time = '11:40',
                    end_time = '14:00',
                    term_id = 3
                WHERE class_id = 8
                """,
                (it119_id,),
            )
            cursor.execute(
                """
                UPDATE classes
                SET course_id = ?,
                    room_id = 11,
                    instructor_id = 1,
                    section = 'AM1',
                    day_of_week = 'Thursday',
                    start_time = '12:50',
                    end_time = '15:10',
                    term_id = 2
                WHERE class_id = 9
                """,
                (it119l_id,),
            )

            terms = [
                (1, "2025-2026", 1, "2025-08-01", "2025-10-31"),
                (2, "2025-2026", 2, "2025-11-01", "2026-01-31"),
                (3, "2025-2026", 3, "2026-02-01", "2026-05-31"),
                (4, "2026-2027", 1, "2026-08-01", "2026-10-31"),
                (5, "2026-2027", 2, "2026-11-01", "2027-01-31"),
                (6, "2026-2027", 3, "2027-02-01", "2027-05-31"),
            ]
            cursor.execute("DELETE FROM school_terms")
            cursor.executemany(
                """
                INSERT INTO school_terms (term_id, school_year, term, start_date, end_date)
                VALUES (?, ?, ?, ?, ?)
                """,
                terms,
            )

            cursor.execute(
                """
                DELETE FROM role_permissions
                WHERE permission_id NOT IN (SELECT permission_id FROM permissions)
                """
            )
            cursor.execute(
                """
                DELETE FROM role_permissions
                WHERE user_type = 'it_admin'
                  AND permission_id IN (
                      SELECT permission_id FROM permissions
                      WHERE permission_name IN ('view_analytics', 'View system analytics and reports')
                  )
                """
            )

            tickets = [
                (
                    1,
                    5,
                    "Attendance report export review",
                    "Requesting verification of the exported attendance report for IT119L-AM1 before submission to the department.",
                    "attendance",
                    "medium",
                    "resolved",
                    11,
                    "2026-01-15 15:35:00",
                    "2026-01-15 16:05:00",
                ),
                (
                    2,
                    8,
                    "Room schedule confirmation",
                    "Please confirm that IT119-AM1 is assigned to Room 604 for the February 5, 2026 session.",
                    "schedule",
                    "low",
                    "closed",
                    11,
                    "2026-02-04 09:20:00",
                    "2026-02-04 10:10:00",
                ),
                (
                    3,
                    1,
                    "Facial recognition registration assistance",
                    "A student requested confirmation that their captured facial data is available for class attendance validation.",
                    "technical",
                    "medium",
                    "resolved",
                    11,
                    "2026-02-05 10:15:00",
                    "2026-02-05 11:00:00",
                ),
                (
                    4,
                    11,
                    "SME testing account preparation",
                    "Prepare and verify the CIS101 testing setup for the scheduled subject matter expert evaluation.",
                    "account",
                    "high",
                    "in_progress",
                    13,
                    "2026-02-06 08:30:00",
                    "2026-02-06 09:00:00",
                ),
            ]
            cursor.executemany(
                """
                INSERT INTO support_tickets (
                    ticket_id, user_id, subject, description, category, priority,
                    status, assigned_to, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                tickets,
            )

            replies = [
                (
                    1,
                    1,
                    11,
                    "The attendance export has been checked against the class roster and is ready for submission.",
                    0,
                    "2026-01-15 16:05:00",
                ),
                (
                    2,
                    2,
                    11,
                    "Room 604 has been confirmed for the February 5 IT119-AM1 session.",
                    0,
                    "2026-02-04 10:10:00",
                ),
                (
                    3,
                    3,
                    11,
                    "The student facial record is available and can be used for attendance recognition.",
                    0,
                    "2026-02-05 11:00:00",
                ),
                (
                    4,
                    4,
                    13,
                    "CIS101 has been added to the course catalog. Final account checks are in progress.",
                    0,
                    "2026-02-06 09:00:00",
                ),
            ]
            cursor.executemany(
                """
                INSERT INTO ticket_replies (
                    reply_id, ticket_id, user_id, message, is_internal, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                replies,
            )

            audit_rows = [
                (
                    1,
                    13,
                    "Updated school terms",
                    "school_terms",
                    1,
                    json.dumps({"school_year": "2025-2026", "terms": 3}),
                    "127.0.0.1",
                    "2026-01-10 09:00:00",
                ),
                (
                    2,
                    11,
                    "Verified attendance records",
                    "attendance_logs",
                    8,
                    json.dumps({"class_id": 8, "attendance_date": "2026-02-05"}),
                    "127.0.0.1",
                    "2026-02-05 15:30:00",
                ),
                (
                    3,
                    13,
                    "Prepared SME testing course",
                    "courses",
                    None,
                    json.dumps({"course_code": "CIS101"}),
                    "127.0.0.1",
                    "2026-02-06 08:45:00",
                ),
            ]
            cursor.executemany(
                """
                INSERT INTO audit_logs (
                    log_id, user_id, action, resource_type, resource_id,
                    details, ip_address, timestamp
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                audit_rows,
            )

            sequence_targets = {
                "courses": len(courses),
                "rooms": 14,
                "school_terms": 6,
                "support_tickets": 4,
                "ticket_replies": 4,
                "audit_logs": 3,
            }
            for table, seq in sequence_targets.items():
                set_sequence(cursor, table, seq)

            cursor.execute("COMMIT")
        except Exception:
            cursor.execute("ROLLBACK")
            raise

        cursor.execute("PRAGMA foreign_keys = ON")

        summary = {}
        for table in [
            "users",
            "instructors",
            "admins",
            "it_admins",
            "super_admins",
            "students",
            "courses",
            "rooms",
            "school_terms",
            "classes",
            "enrollments",
            "attendance_logs",
            "support_tickets",
            "ticket_replies",
            "audit_logs",
            "student_face_embeddings",
        ]:
            cursor.execute(f"SELECT COUNT(*) FROM {table}")
            summary[table] = cursor.fetchone()[0]

        return summary


if __name__ == "__main__":
    result = cleanup_database()
    for table_name, count in result.items():
        print(f"{table_name}: {count}")
